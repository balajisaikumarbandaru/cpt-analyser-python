"""CSV / TXT and Excel input.

Columns are found by header name, or taken in the CPTask Excel order
length, qc, fs, u, I, Qst, Qt when there is no header row.
"""
from __future__ import annotations

import datetime as _dt
import re

from .._js import NAN, base_name, fin, num, opt_num
from ..models import CPT, Header, RawData


def _find(H, pattern, not_pattern=None):
    rx = re.compile(pattern, re.ASCII)
    nx = re.compile(not_pattern, re.ASCII) if not_pattern else None
    return next((i for i, h in enumerate(H) if rx.search(h) and not (nx and nx.search(h))), -1)


def map_header_cols(hdr):
    H = [("" if c is None else str(c)).lower() for c in hdr]
    m = {
        "L": _find(H, r"depth|length|penetr|^z\b"),
        "qc": _find(H, r"\bqc\b|cone res", r"qt|net|corr"),
        "fs": _find(H, r"\bfs\b|sleeve|local fric", r"ratio|rf"),
        "u1": _find(H, r"\bu1\b"), "u3": _find(H, r"\bu3\b"),
        "u2": _find(H, r"\bu2?\b|pore|dynamic", r"u0|u1|u3|excess|ratio|/qc"),
        "I": _find(H, r"incl"),
        "Qst": _find(H, r"qst|total fric"), "Qtot": _find(H, r"total force|\bqt\b.*kn|qtot"),
    }

    def kpa(k):
        return 1e-3 if m[k] >= 0 and "kpa" in H[m[k]] else 1

    return m, {k: kpa(k) for k in ("qc", "fs", "u1", "u2", "u3")}


def _cell(r, i):
    return r[i] if 0 <= i < len(r) else None


def rows_to_cpt(rows, fname, h: Header, source) -> CPT:
    def is_text(c):
        return isinstance(c, str) and re.search(r"[a-z]", c, re.I) and not fin(num(c))

    hi = next((i for i, r in enumerate(rows) if any(is_text(c) for c in r)), -1)
    first_num = next((i for i, r in enumerate(rows) if fin(num(_cell(r, 0))) and fin(num(_cell(r, 1)))), -1)
    if hi > first_num:
        hi = -1
    m = {"L": 0, "qc": 1, "fs": 2, "u2": 3, "I": 4, "Qst": 5, "Qtot": 6, "u1": -1, "u3": -1}
    sc = {"qc": 1, "fs": 1, "u1": 1, "u2": 1, "u3": 1}
    if hi >= 0:
        hm, sc = map_header_cols(rows[hi])
        for k, v in hm.items():
            if v >= 0:
                m[k] = v
            elif k not in ("L", "qc"):
                m[k] = -1
    raw = RawData()
    start = max(hi + 1, first_num)
    nc = max((len(r) for r in rows[start:]), default=0)
    if m["Qst"] >= nc:
        m["Qst"] = -1
    if m["Qtot"] >= nc:
        m["Qtot"] = -1
    if m["Qst"] >= 0:
        raw.x["Qst"] = []
    if m["Qtot"] >= 0:
        raw.x["Qtot"] = []
    for r in rows[start:]:
        L = num(_cell(r, m["L"]))
        if not fin(L):
            continue
        raw.L.append(L)
        raw.qc.append(num(_cell(r, m["qc"])) * sc["qc"])
        for k in ("fs", "u1", "u2", "u3"):
            getattr(raw, k).append(num(_cell(r, m[k])) * sc[k] if m[k] >= 0 else NAN)
        raw.I.append(num(_cell(r, m["I"])) if m["I"] >= 0 else NAN)
        raw.t.append(NAN)
        if m["Qst"] >= 0:
            raw.x["Qst"].append(num(_cell(r, m["Qst"])))
        if m["Qtot"] >= 0:
            raw.x["Qtot"].append(num(_cell(r, m["Qtot"])))
    if not raw.L:
        raise ValueError("No numeric depth / qc rows found")
    for k in ("Qst", "Qtot"):
        if k in raw.x and not any(fin(v) for v in raw.x[k]):
            del raw.x[k]
    if not h.test_id:
        h.test_id = base_name(fname)
    return CPT(fname, h, raw, source)


def parse_delimited(text: str, fname: str = "file.csv") -> CPT:
    lines = [ln for ln in re.split(r"\r?\n", text) if ln.strip()]
    sample = "\n".join(lines[:30])
    n_tab, n_semi, n_comma = sample.count("\t"), sample.count(";"), sample.count(",")
    # same non-zero number of ';' on (nearly) every line → ';' separated, even with decimal commas (1,25;3,40;…)
    semis = [ln.count(";") for ln in lines[:30]]
    even_semi = any(c > 0 and semis.count(c) >= 0.9 * len(semis) for c in semis)
    if n_tab > 20:
        split = lambda s: s.split("\t")  # noqa: E731
    elif even_semi or n_semi > n_comma:
        split = lambda s: s.split(";")  # noqa: E731
    elif n_comma > 20:
        split = lambda s: s.split(",")  # noqa: E731
    else:
        split = lambda s: re.split(r"\s+", s)  # noqa: E731
    rows = [[re.sub(r'^"|"$', "", c.strip()) for c in split(ln)] for ln in lines]
    return rows_to_cpt(rows, fname, Header(), "csv")


def parse_xlsx(path_or_bytes, fname: str | None = None) -> CPT:
    """Excel file with Header + Data sheets (CPTask style) or any sheet with CPT columns."""
    import io

    from openpyxl import load_workbook
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes
    fname = fname or str(path_or_bytes)
    wb = load_workbook(src, read_only=True, data_only=True)
    names = wb.sheetnames
    dn = next((n for n in names if re.search("data", n, re.I)), names[0])
    rows = [["" if v is None else v for v in r] for r in wb[dn].iter_rows(values_only=True)]
    h = Header()
    hn = next((n for n in names if re.search("header", n, re.I)), None)
    if hn:
        for r in wb[hn].iter_rows(values_only=True):
            if not r:
                continue
            k = r[0]
            v = r[1] if len(r) > 1 else None
            v = "" if v is None else v
            key = str("" if k is None else k).lower()
            if re.search("sounding|test id|cpt no", key):
                h.test_id = str(v)
            elif "client" in key:
                h.project = str(v)
            elif "job" in key:
                h.job_no = str(v)
            elif "location" in key:
                h.location = str(v)
            elif "cone" in key:
                h.cone_type = str(v)
            elif key.startswith("date"):
                h.date = v.isoformat()[:10] if isinstance(v, (_dt.date, _dt.datetime)) else str(v)
            elif "water" in key:
                h.gwl = abs(opt_num(v)) if opt_num(v) is not None else None
            elif "ground level" in key:
                h.gl = opt_num(v)
            elif re.search("predrill|pre-drill", key):
                h.predrill = opt_num(v) or 0
            elif re.search("^e coord|easting", key):
                h.x = opt_num(v)
            elif re.search("^n coord|northing", key):
                h.y = opt_num(v)
            elif "alpha" in key:
                h.alpha = opt_num(v)
            elif "beta" in key:
                h.beta = opt_num(v)
            elif "remark" in key:
                h.remarks = str(v)
    wb.close()
    return rows_to_cpt(rows, fname, h, "xlsx")
